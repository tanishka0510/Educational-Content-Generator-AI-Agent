"""
Office Document Loader

Supports:
- DOCX
- PPTX
- PPT
"""

import re
from pathlib import Path


class OfficeLoader:

    @staticmethod
    def load(file_path: str) -> str:
        lower_path = file_path.lower()

        # --------------------------------------------------
        # DOCX
        # --------------------------------------------------
        if lower_path.endswith(".docx"):
            try:
                import docx
                doc = docx.Document(file_path)
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_text:
                            paragraphs.append(row_text)
                if paragraphs:
                    return "\n\n".join(paragraphs)
            except Exception as e:
                print(f"python-docx loader fallback: {e}")

            from langchain_community.document_loaders import UnstructuredWordDocumentLoader
            loader = UnstructuredWordDocumentLoader(file_path)
            documents = loader.load()
            return "\n".join(doc.page_content for doc in documents)

        # --------------------------------------------------
        # PPTX (Modern PowerPoint)
        # --------------------------------------------------
        elif lower_path.endswith(".pptx"):
            try:
                from pptx import Presentation
                prs = Presentation(file_path)
                extracted_lines = []
                for slide_idx, slide in enumerate(prs.slides, start=1):
                    slide_texts = []
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for paragraph in shape.text_frame.paragraphs:
                                clean_p = paragraph.text.strip()
                                if clean_p:
                                    slide_texts.append(clean_p)
                        elif shape.has_table:
                            for row in shape.table.rows:
                                row_str = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                                if row_str:
                                    slide_texts.append(row_str)
                    if slide_texts:
                        extracted_lines.append(f"--- Slide {slide_idx} ---\n" + "\n".join(slide_texts))
                if extracted_lines:
                    return "\n\n".join(extracted_lines)
            except Exception as e:
                print(f"python-pptx extraction failed: {e}, attempting fallback loader")

            from langchain_community.document_loaders import UnstructuredPowerPointLoader
            loader = UnstructuredPowerPointLoader(file_path)
            documents = loader.load()
            return "\n".join(doc.page_content for doc in documents)

        # --------------------------------------------------
        # PPT (Legacy PowerPoint 97-2003)
        # --------------------------------------------------
        elif lower_path.endswith(".ppt"):
            try:
                from langchain_community.document_loaders import UnstructuredPowerPointLoader
                loader = UnstructuredPowerPointLoader(file_path)
                documents = loader.load()
                if documents:
                    return "\n".join(doc.page_content for doc in documents)
            except Exception as e:
                print(f"UnstructuredPowerPointLoader for .ppt failed: {e}")

            # Fallback binary text extraction via olefile
            try:
                import olefile
                if olefile.isOleFile(file_path):
                    ole = olefile.OleFileIO(file_path)
                    if ole.exists('PowerPoint Document'):
                        stream = ole.openstream('PowerPoint Document').read()
                        matches = re.findall(rb'[\x20-\x7E\t\n]{4,}', stream)
                        text = "\n".join(m.decode('latin-1', errors='ignore') for m in matches)
                        if len(text.strip()) > 50:
                            return text
            except Exception as e:
                print(f"OLE extraction for .ppt failed: {e}")

            raise ValueError(
                f"Could not extract content from legacy PowerPoint file: {Path(file_path).name}. "
                "Please save the presentation as .pptx format and upload again."
            )

        else:
            raise ValueError(
                f"Unsupported Office file format: {file_path}"
            )